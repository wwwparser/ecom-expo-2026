# -*- coding: utf-8 -*-
"""
Сбор участников ВСЕХ архивных выставок ECOM Expo (expo.oborot.ru/archive/).

Два формата страниц:
  • Современные годы (≈2020–2025): архив/<год>/ — блоки .exh_title (как на главной).
  • Старые годы (≈2012–2019): архив/<год>/plan.html — плоский список ссылок
    вида  <a href="сайт">Название</a> (Стенд) (краткое описание).  Кодировка cp1251.

Результат:
  archive_participants.jsonl  — по записи на (год, компания)
  archive_participants.xlsx   — лист на каждый год + лист «Сводка»
  archive_unique_sites.txt    — уникальные домены (для сбора контактов)
"""
from __future__ import annotations
import io, json, re, sys
from pathlib import Path
from urllib.parse import urljoin, urlparse
import httpx
from bs4 import BeautifulSoup

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

YEARS = list(range(2025, 2011, -1))
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/124 Safari/537.36"}
BOOTH_RE = re.compile(r"\(([A-ZА-Я]{0,2}\d[\w.\-,]*)\)")  # (E9), (P5.8), (B4.1,...)
OWN_HOST = "oborot.ru"


def fetch(url: str) -> str | None:
    try:
        r = httpx.get(url, headers=UA, follow_redirects=True, timeout=30)
        if r.status_code != 200:
            return None
        raw = r.content
        head = raw[:1500].lower()
        if b"charset=windows-1251" in head or b"charset=cp1251" in head:
            return raw.decode("cp1251", "ignore")
        # пробуем utf-8; если много � — значит cp1251
        txt = raw.decode("utf-8", "replace")
        if txt.count("�") > 50:
            return raw.decode("cp1251", "ignore")
        return txt
    except Exception as e:
        print(f"  fetch error {url}: {e}", file=sys.stderr)
        return None


def is_external(href: str) -> bool:
    if not href.startswith("http"):
        return False
    host = urlparse(href).netloc.lower()
    return host and OWN_HOST not in host


def parse_modern(soup: BeautifulSoup, base: str) -> list:
    """Годы с .exh_title."""
    out = []
    for t in soup.select(".exh_title"):
        a = t.find("a", href=True)
        name = (a.get_text(" ", strip=True) if a else t.get_text(" ", strip=True)).strip()
        if not name:
            continue
        site = a["href"].strip() if a and is_external(a["href"]) else ""
        ttext = t.get_text(" ", strip=True)
        bm = BOOTH_RE.search(ttext)
        booth = bm.group(1) if bm else ""
        descr = ""
        sib = t.find_next_sibling(class_="exh_descr") or \
            (t.parent.select_one(".exh_descr") if t.parent else None)
        if sib:
            descr = sib.get_text(" ", strip=True)
        out.append({"name": name, "site": site, "booth": booth, "descr": descr})
    return out


def parse_legacy(soup: BeautifulSoup, base: str) -> list:
    """plan.html: <a href=сайт>Название</a> (Стенд) (описание)."""
    out = []
    seen = set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not is_external(href):
            continue
        name = a.get_text(" ", strip=True)
        if not name or len(name) > 80:
            continue
        # хвостовой текст после ссылки: "(E9) (описание)"
        tail = ""
        for sib in a.next_siblings:
            s = sib.get_text(" ", strip=True) if hasattr(sib, "get_text") else str(sib).strip()
            if s:
                tail = s
                break
        booth, descr = "", ""
        if tail:
            bm = BOOTH_RE.search(tail)
            if bm:
                booth = bm.group(1)
            paren = re.findall(r"\(([^)]*)\)", tail)
            # описание — последняя скобка, если это не стенд
            descrs = [p for p in paren if not BOOTH_RE.fullmatch(f"({p})")]
            if descrs:
                descr = descrs[-1].strip()
        key = (name.lower(), urlparse(href).netloc.lower())
        if key in seen:
            continue
        seen.add(key)
        out.append({"name": name, "site": href, "booth": booth, "descr": descr})
    return out


def scrape_year(year: int) -> list:
    base_main = f"https://expo.oborot.ru/archive/{year}/"
    html = fetch(base_main)
    items = []
    fmt = ""
    if html:
        soup = BeautifulSoup(html, "lxml")
        if soup.select(".exh_title"):
            items = parse_modern(soup, base_main)
            fmt = "modern"
    if not items:
        plan = fetch(base_main + "plan.html")
        if plan:
            # старые plan.html — невалидная вёрстка, lxml её обрезает → html.parser
            soup = BeautifulSoup(plan, "html.parser")
            items = parse_legacy(soup, base_main)
            fmt = "legacy(plan.html)"
    for it in items:
        it["year"] = year
    print(f"{year}: {len(items):>4} участников  [{fmt or 'нет данных'}]")
    return items


def main():
    all_items = []
    for y in YEARS:
        all_items += scrape_year(y)

    out = Path(".")
    with (out / "archive_participants.jsonl").open("w", encoding="utf-8") as f:
        for it in all_items:
            f.write(json.dumps(it, ensure_ascii=False) + "\n")

    # Excel: лист на год
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    summ = wb.active
    summ.title = "Сводка"
    summ.append(["Год", "Участников"])
    for c in summ[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F7A3D")
    by_year = {}
    for it in all_items:
        by_year.setdefault(it["year"], []).append(it)
    for y in sorted(by_year, reverse=True):
        summ.append([y, len(by_year[y])])
        ws = wb.create_sheet(str(y))
        ws.append(["Компания", "Сайт", "Стенд", "Краткое описание"])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F7A3D")
        for it in by_year[y]:
            ws.append([it["name"], it["site"], it["booth"], it["descr"]])
        for i, w in enumerate([34, 36, 12, 50], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    summ.column_dimensions["A"].width = 10
    summ.column_dimensions["B"].width = 14
    wb.save(out / "archive_participants.xlsx")

    # уникальные домены
    domains = {}
    for it in all_items:
        if it["site"]:
            d = urlparse(it["site"]).netloc.lower()
            if d and d not in domains:
                domains[d] = it["site"]
    (out / "archive_unique_sites.txt").write_text(
        "\n".join(sorted(domains.values())), encoding="utf-8")

    print(f"\nИТОГО: {len(all_items)} строк, уникальных доменов: {len(domains)}")
    print("→ archive_participants.xlsx / .jsonl / archive_unique_sites.txt")


if __name__ == "__main__":
    main()
