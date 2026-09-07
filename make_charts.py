# -*- coding: utf-8 -*-
"""
Рендер статистики ECOM Expo в PNG-картинки для статьи на vc.ru.
Светлый фон, крупный шрифт, фирменный зелёный акцент.

Источники: archive_participants.jsonl, exhibitors_ecom_expo_2026.xlsx,
archive_topics.jsonl, contacts_out/ + archive_contacts_out/.
Выход: папка vc_article/*.png
"""
from __future__ import annotations
import io, json, sys
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = Path(__file__).parent
OUT = ROOT / "vc_article"
OUT.mkdir(exist_ok=True)

plt.rcParams["font.family"] = "DejaVu Sans"
plt.rcParams["axes.unicode_minus"] = False
GREEN = "#1f7a3d"
GREEN2 = "#27ae60"
PURPLE = "#7d4fb0"
BLUE = "#2b7aa0"
INK = "#1a1f27"
MUT = "#7b8794"
SRC = "Данные: expo.oborot.ru/archive · 2015–2026 · сбор и анализ — автоматический"


def norm(url):
    if not url:
        return ""
    if not url.startswith("http"):
        url = "http://" + url
    h = urlparse(url).netloc.lower()
    return h[4:] if h.startswith("www.") else h


def load():
    years = defaultdict(set)
    name = {}
    for l in (ROOT / "archive_participants.jsonl").read_text(encoding="utf-8").splitlines():
        r = json.loads(l)
        if not r["site"]:
            continue
        d = norm(r["site"])
        years[d].add(r["year"])
        name.setdefault(d, r["name"] or d)
        if r["name"]:
            name[d] = r["name"]
    # 2026 из xlsx
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "exhibitors_ecom_expo_2026.xlsx", read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title == "Сводка":
            continue
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if "Сайт" not in hdr:
            continue
        si, ni = hdr.index("Сайт"), hdr.index("Компания")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if si < len(row) and row[si]:
                d = norm(str(row[si]))
                if d:
                    years[d].add(2026)
                    name.setdefault(d, str(row[ni]) if ni < len(row) else d)
    topics = {}
    tf = ROOT / "archive_topics.jsonl"
    if tf.exists():
        for l in tf.read_text(encoding="utf-8").splitlines():
            r = json.loads(l)
            topics[r["domain"]] = r["topic"]
    contacts = {}
    for jf in [ROOT / "archive_contacts_out/contacts.jsonl", ROOT / "contacts_out/contacts.jsonl"]:
        if jf.exists():
            for l in jf.read_text(encoding="utf-8").splitlines():
                try:
                    r = json.loads(l)
                except Exception:
                    continue
                d = norm(r.get("home") or r.get("site", ""))
                if d and (d not in contacts or (r.get("telegram") and not contacts[d].get("telegram"))):
                    contacts[d] = r
    return years, name, topics, contacts


def style(ax, title, subtitle=None):
    pad = 34 if subtitle else 14
    ax.set_title(title, fontsize=18, fontweight="bold", color=INK, pad=pad, loc="left")
    if subtitle:
        ax.text(0, 1.015, subtitle, transform=ax.transAxes, fontsize=11, color=MUT)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#d7dde4")
    ax.spines["bottom"].set_color("#d7dde4")
    ax.tick_params(colors=MUT, labelsize=12)
    ax.grid(axis="y", color="#eef1f4", linewidth=1)
    ax.set_axisbelow(True)


def footer(fig):
    fig.text(0.012, 0.012, SRC, fontsize=9, color=MUT)


def save(fig, name):
    fig.tight_layout(rect=[0, 0.03, 1, 1])
    p = OUT / name
    fig.savefig(p, dpi=140, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    print("  ", p.name)


def main():
    years, name, topics, contacts = load()
    ally = sorted({y for s in years.values() for y in s})

    # 1. участников по годам
    per = {y: sum(1 for s in years.values() if y in s) for y in ally}
    fig, ax = plt.subplots(figsize=(11, 6))
    bars = ax.bar([str(y) for y in ally], [per[y] for y in ally], color=GREEN2, width=0.7)
    for b, y in zip(bars, ally):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 8, str(per[y]),
                ha="center", fontsize=11, color=INK, fontweight="bold")
    style(ax, "Сколько компаний участвовало в ECOM Expo по годам",
          "2018–2019 — данных в открытом доступе нет")
    ax.set_ylim(0, max(per.values()) * 1.15)
    footer(fig); save(fig, "01_participants_by_year.png")

    # 2. новые компании по годам
    first = {d: min(s) for d, s in years.items()}
    newp = {y: sum(1 for f in first.values() if f == y) for y in ally}
    fig, ax = plt.subplots(figsize=(11, 6))
    bars = ax.bar([str(y) for y in ally], [newp[y] for y in ally], color=BLUE, width=0.7)
    for b, y in zip(bars, ally):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 4, str(newp[y]),
                ha="center", fontsize=11, color=INK, fontweight="bold")
    style(ax, "Новые компании по годам (первое участие)",
          "Год, когда компания впервые появилась на выставке")
    ax.set_ylim(0, max(newp.values()) * 1.15)
    footer(fig); save(fig, "02_new_companies_by_year.png")

    # 3. топ-20 по участиям
    ranked = sorted(years.items(), key=lambda x: (-len(x[1]), name.get(x[0], x[0]).lower()))[:20]
    labels = [name.get(d, d)[:28] for d, _ in ranked][::-1]
    vals = [len(s) for _, s in ranked][::-1]
    fig, ax = plt.subplots(figsize=(11, 9))
    bars = ax.barh(labels, vals, color=GREEN, height=0.72)
    for b, v in zip(bars, vals):
        ax.text(b.get_width() + 0.1, b.get_y() + b.get_height() / 2, str(v),
                va="center", fontsize=11, color=INK, fontweight="bold")
    style(ax, "Топ-20 компаний по числу участий в ECOM Expo",
          "Сколько разных выставок (лет) компания участвовала, 2015–2026")
    ax.set_xlim(0, max(vals) * 1.1)
    ax.grid(axis="y", linewidth=0)
    ax.grid(axis="x", color="#eef1f4")
    footer(fig); save(fig, "03_top_companies.png")

    # 4. тематики
    if topics:
        tc = Counter(topics.get(d, "Прочее") for d in years)
        items = tc.most_common()[::-1]
        labels = [t for t, _ in items]
        vals = [n for _, n in items]
        fig, ax = plt.subplots(figsize=(11, 6.5))
        bars = ax.barh(labels, vals, color=PURPLE, height=0.7)
        for b, v in zip(bars, vals):
            ax.text(b.get_width() + 2, b.get_y() + b.get_height() / 2, str(v),
                    va="center", fontsize=11, color=INK, fontweight="bold")
        style(ax, "Участники ECOM Expo по тематикам", "Классификация по описанию деятельности (1222 компании)")
        ax.set_xlim(0, max(vals) * 1.12)
        ax.grid(axis="y", linewidth=0); ax.grid(axis="x", color="#eef1f4")
        footer(fig); save(fig, "04_topics.png")

    # 5. рост мессенджеров
    tg_share, mx_share, yrs = [], [], []
    for y in ally:
        doms = [d for d, s in years.items() if y in s]
        if not doms:
            continue
        yrs.append(str(y))
        tg_share.append(round(100 * sum(1 for d in doms if contacts.get(d, {}).get("telegram")) / len(doms)))
        mx_share.append(round(100 * sum(1 for d in doms if contacts.get(d, {}).get("max")) / len(doms)))
    fig, ax = plt.subplots(figsize=(11, 6))
    ax.plot(yrs, tg_share, "-o", color=BLUE, linewidth=2.5, markersize=7, label="Telegram")
    ax.plot(yrs, mx_share, "-o", color=PURPLE, linewidth=2.5, markersize=7, label="MAX")
    for x, v in zip(yrs, tg_share):
        ax.text(x, v + 2, f"{v}%", ha="center", fontsize=10, color=BLUE, fontweight="bold")
    style(ax, "У какой доли участников каждого года сейчас есть Telegram / MAX",
          "Текущий срез контактов по сайтам компаний")
    ax.legend(fontsize=12, frameon=False)
    ax.set_ylim(0, max(tg_share) * 1.25)
    ax.set_ylabel("% компаний", color=MUT)
    footer(fig); save(fig, "05_messengers_by_year.png")

    # 6. распределение по стажу
    streak = Counter(len(s) for s in years.values())
    ks = sorted(streak)
    fig, ax = plt.subplots(figsize=(11, 6))
    bars = ax.bar([str(k) for k in ks], [streak[k] for k in ks], color="#e08a2b", width=0.7)
    for b, k in zip(bars, ks):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 6, str(streak[k]),
                ha="center", fontsize=11, color=INK, fontweight="bold")
    style(ax, "Сколько раз компании участвовали в выставке",
          "По горизонтали — число участий, по вертикали — сколько таких компаний")
    ax.set_xlabel("число участий (лет)", color=MUT)
    ax.set_ylim(0, max(streak.values()) * 1.15)
    footer(fig); save(fig, "06_tenure.png")

    # 7. таблица топ-25 (картинкой)
    top = sorted(years.items(), key=lambda x: (-len(x[1]), name.get(x[0], x[0]).lower()))[:25]
    rows = [[str(i), name.get(d, d)[:34], str(len(s)),
             ", ".join(str(y) for y in sorted(s, reverse=True))]
            for i, (d, s) in enumerate(top, 1)]
    fig, ax = plt.subplots(figsize=(12, 11))
    ax.axis("off")
    ax.set_title("Топ-25 компаний по числу участий в ECOM Expo (2015–2026)",
                 fontsize=18, fontweight="bold", color=INK, loc="left", pad=18)
    tbl = ax.table(cellText=rows, colLabels=["#", "Компания", "Участий", "Годы"],
                   colWidths=[0.05, 0.30, 0.09, 0.56], cellLoc="left", loc="upper left")
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(11)
    tbl.scale(1, 1.6)
    for (r, c), cell in tbl.get_celld().items():
        cell.set_edgecolor("#e3e8ee")
        if r == 0:
            cell.set_facecolor(GREEN)
            cell.set_text_props(color="white", fontweight="bold")
        else:
            cell.set_facecolor("#ffffff" if r % 2 else "#f5f8fa")
            if c == 2:
                cell.set_text_props(color=GREEN, fontweight="bold")
    fig.text(0.012, 0.01, SRC, fontsize=9, color=MUT)
    fig.savefig(OUT / "07_top25_table.png", dpi=140, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    print("   07_top25_table.png")

    # 8. обложка статьи 1920x1080 (16:9), СВЕТЛЫЙ фон — по рекомендациям vc.ru
    n_companies = len(years)
    yy = sorted(ally); maxp = max(per.values())
    fig = plt.figure(figsize=(19.2, 10.8), dpi=100)
    fig.patch.set_facecolor("#ffffff")
    # нижняя зона — спокойный график-столбики как «герой»
    axc = fig.add_axes([0.06, 0.06, 0.88, 0.40]); axc.axis("off")
    bars = axc.bar([str(y) for y in yy], [per[y] for y in yy], color="#27ae60", width=0.72)
    for b, y in zip(bars, yy):
        axc.text(b.get_x() + b.get_width() / 2, b.get_height() + maxp * 0.03, str(per[y]),
                 ha="center", fontsize=13, color="#5b6670")
        axc.text(b.get_x() + b.get_width() / 2, -maxp * 0.06, str(y),
                 ha="center", fontsize=12, color="#9aa4ae")
    axc.set_ylim(0, maxp * 1.18)
    # текстовый блок сверху на белом
    ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
    ax.add_patch(plt.Rectangle((0.06, 0.86), 0.10, 0.012, color="#27ae60", transform=ax.transAxes))
    ax.text(0.06, 0.74, "11 лет ECOM Expo", fontsize=58, fontweight="bold",
            color="#16202b", transform=ax.transAxes)
    ax.text(0.06, 0.635, "Кто приезжает каждый год и куда движется\nрынок e-commerce-сервисов",
            fontsize=27, color="#46535f", transform=ax.transAxes, va="top")
    ax.text(0.06, 0.50, f"{n_companies} компаний · {min(ally)}–{max(ally)} · анализ архива выставки",
            fontsize=18, color="#1f7a3d", transform=ax.transAxes, fontweight="bold")
    fig.savefig(OUT / "00_cover.png", facecolor="#ffffff")
    plt.close(fig)
    print("   00_cover.png")

    print(f"\nГотово. Картинки в {OUT}")


if __name__ == "__main__":
    main()
